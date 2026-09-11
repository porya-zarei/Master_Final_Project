"""Post-process a generated .docx for Persian RTL + B Nazanin font."""
import sys
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "B Nazanin"
path = sys.argv[1]


def rtl_para(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:bidi')) is None:
        pPr.append(OxmlElement('w:bidi'))
    # RTL text alignment (right)
    try:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    except Exception:
        pass


def rtl_run(r):
    rPr = r._r.get_or_add_rPr()
    if rPr.find(qn('w:rtl')) is None:
        rPr.append(OxmlElement('w:rtl'))
    r.font.name = FONT
    r.font.size = r.font.size  # keep size
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)


def walk(doc):
    # body paragraphs
    for p in doc.paragraphs:
        rtl_para(p)
        for r in p.runs:
            rtl_run(r)
    # tables
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    rtl_para(p)
                    for r in p.runs:
                        rtl_run(r)
    # headers/footers
    for section in doc.sections:
        for hp in section.header.paragraphs + section.footer.paragraphs:
            rtl_para(hp)
            for r in hp.runs:
                rtl_run(r)


# also set the Normal style font
doc = Document(path)
st = doc.styles['Normal']
st.font.name = FONT
st.font.size = None

walk(doc)
doc.save(path)
print("RTL+font applied to", path)
